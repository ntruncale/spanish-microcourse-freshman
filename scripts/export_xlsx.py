"""Make teacher/blackbaud-import/all-quizzes.xlsx from the CSV files that scripts/teacher.js writes.

One sheet with every question, then one sheet per module.
Usage: python3 scripts/export_xlsx.py
"""
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "teacher" / "blackbaud-import"
WIDTHS = [8, 34, 10, 16, 7, 9, 60, 26, 26, 26, 26, 10, 26, 34, 40]


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def fill(ws, rows):
    for r in rows:
        ws.append([int(c) if c.isdigit() else c for c in r])
    head = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = head
        cell.fill = PatternFill("solid", fgColor="AA151B")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "All quizzes"
    fill(ws, read(OUT / "all-quizzes.csv"))
    for f in sorted((OUT / "csv").glob("m*.csv")):
        fill(wb.create_sheet(f.stem.upper()), read(f))
    wb.save(OUT / "all-quizzes.xlsx")
    print("wrote", OUT / "all-quizzes.xlsx")


if __name__ == "__main__":
    main()
